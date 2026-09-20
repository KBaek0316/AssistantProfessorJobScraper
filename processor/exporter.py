import csv
import logging
from typing import List
from scrapers.base import JobPosting


class JobExporter:
    """Exports job postings to CSV and formatted Excel (.xlsx) files with Fit Score highlighting."""

    COLUMNS = [
        ("title", "Title"),
        ("institution", "Institution"),
        ("field", "Field/Division"),
        ("fit_score", "Fit Score (1-10)"),
        ("fit_reason", "Fit Reason"),
        ("research_topics", "Research Topics"),
        ("tenure_track", "Tenure Track"),
        ("location", "Location"),
        ("deadline", "Deadline"),
        ("deadline_date", "Deadline Date"),
        ("salary", "Salary"),
        ("summary", "Summary (Gemini)"),
        ("link", "Link"),
        ("source", "Source"),
        ("date_first_seen", "Date Added"),
        ("date_last_verified", "Last Verified"),
        ("status", "Status"),
        ("latitude", "Latitude"),
        ("longitude", "Longitude"),
        ("id", "ID"),
    ]

    def __init__(self, csv_filepath: str = "jobs.csv", excel_filepath: str = "jobs.xlsx"):
        self.csv_filepath = csv_filepath
        self.excel_filepath = excel_filepath
        self.logger = logging.getLogger("processor.exporter")

    def export_csv(self, postings: List[JobPosting], include_filtered: bool = False):
        """Export postings to CSV file."""
        if not include_filtered:
            postings = [p for p in postings if not p.status.startswith("Filtered")]

        from scrapers.base import sort_postings_by_deadline
        postings = sort_postings_by_deadline(postings)

        fieldnames = [key for key, _ in self.COLUMNS]
        header_labels = [label for _, label in self.COLUMNS]

        try:
            with open(self.csv_filepath, mode="w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerow(header_labels)
                for p in postings:
                    p_dict = p.to_dict()
                    row = [p_dict.get(key, "") if p_dict.get(key) is not None else "" for key in fieldnames]
                    writer.writerow(row)
            self.logger.info(f"Successfully saved {len(postings)} active jobs to {self.csv_filepath}")
        except Exception as e:
            self.logger.error(f"Error exporting CSV to {self.csv_filepath}: {e}", exc_info=True)

    def export_excel(self, postings: List[JobPosting], include_filtered: bool = False):
        """Export postings to an elegantly formatted Excel spreadsheet."""
        if not include_filtered:
            postings = [p for p in postings if not p.status.startswith("Filtered")]

        from scrapers.base import sort_postings_by_deadline
        postings = sort_postings_by_deadline(postings)

        try:
            import pandas as pd
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter

            # Prepare data frame
            fieldnames = [key for key, _ in self.COLUMNS]
            header_labels = [label for _, label in self.COLUMNS]

            rows = []
            for p in postings:
                p_dict = p.to_dict()
                rows.append([p_dict.get(k, "") if p_dict.get(k) is not None else "" for k in fieldnames])

            df = pd.DataFrame(rows, columns=header_labels)

            # Write using openpyxl for formatting
            with pd.ExcelWriter(self.excel_filepath, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="Assistant Professor Jobs")
                worksheet = writer.sheets["Assistant Professor Jobs"]

                # Header formatting: Navy Blue header with white bold text
                header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
                header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

                thin_border = Border(
                    left=Side(style="thin", color="D9D9D9"),
                    right=Side(style="thin", color="D9D9D9"),
                    top=Side(style="thin", color="D9D9D9"),
                    bottom=Side(style="thin", color="D9D9D9"),
                )

                # Fit Score Highlight Fills
                high_fit_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
                high_fit_font = Font(name="Calibri", size=10, bold=True, color="276A3C")
                med_fit_fill = PatternFill(start_color="EDF2F8", end_color="EDF2F8", fill_type="solid")
                med_fit_font = Font(name="Calibri", size=10, bold=False, color="1F497D")

                # Format header row
                for col_num in range(1, len(header_labels) + 1):
                    cell = worksheet.cell(row=1, column=col_num)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = header_align
                worksheet.row_dimensions[1].height = 28

                # Data rows formatting
                data_font = Font(name="Calibri", size=10)
                data_align = Alignment(vertical="center")
                wrapped_align = Alignment(vertical="center", wrap_text=True)
                center_align = Alignment(horizontal="center", vertical="center")

                for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=len(rows) + 1), start=2):
                    worksheet.row_dimensions[row_idx].height = 26
                    for col_idx, cell in enumerate(row, start=1):
                        cell.font = data_font
                        cell.border = thin_border
                        col_name = header_labels[col_idx - 1]

                        if col_name in ("Summary (Gemini)", "Fit Reason"):
                            cell.alignment = wrapped_align
                        elif col_name in ("Fit Score (1-10)", "Tenure Track", "Deadline", "Deadline Date", "Source", "Status"):
                            cell.alignment = center_align
                        else:
                            cell.alignment = data_align

                        # Highlight Fit Score
                        if col_name == "Fit Score (1-10)" and cell.value:
                            try:
                                score = int(cell.value)
                                if score >= 8:
                                    cell.fill = high_fit_fill
                                    cell.font = high_fit_font
                                elif score >= 5:
                                    cell.fill = med_fit_fill
                                    cell.font = med_fit_font
                            except (ValueError, TypeError):
                                pass

                        # Highlight Status
                        if col_name == "Status" and cell.value:
                            val_str = str(cell.value).strip()
                            if val_str == "Past Due":
                                cell.fill = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")
                                cell.font = Font(name="Calibri", size=10, bold=True, color="C5221F")
                            elif val_str == "Active":
                                cell.font = Font(name="Calibri", size=10, bold=True, color="137333")

                        # Add hyperlink to link column
                        if col_name == "Link" and cell.value and str(cell.value).startswith("http"):
                            cell.hyperlink = str(cell.value)
                            cell.font = Font(name="Calibri", size=10, color="0563C1", underline="single")

                # Auto-fit column widths
                for col in worksheet.columns:
                    col_letter = get_column_letter(col[0].column)
                    col_name = col[0].value
                    if col_name == "Summary (Gemini)":
                        worksheet.column_dimensions[col_letter].width = 42
                    elif col_name == "Fit Reason":
                        worksheet.column_dimensions[col_letter].width = 32
                    elif col_name == "Fit Score (1-10)":
                        worksheet.column_dimensions[col_letter].width = 16
                    elif col_name in ("Title", "Institution"):
                        worksheet.column_dimensions[col_letter].width = 30
                    elif col_name == "Link":
                        worksheet.column_dimensions[col_letter].width = 25
                    else:
                        max_len = max(len(str(cell.value or "")) for cell in col)
                        worksheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

            self.logger.info(f"Successfully saved formatted Excel file to {self.excel_filepath}")

        except ImportError:
            self.logger.warning("pandas or openpyxl not installed. Excel export skipped.")
        except Exception as e:
            self.logger.error(f"Error exporting Excel file: {e}", exc_info=True)
